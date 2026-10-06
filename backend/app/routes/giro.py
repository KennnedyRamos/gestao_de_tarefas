from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.core.auth import require_any_permission, require_permission
from app.database.deps import get_db
from app.models.giro import (
    GiroEquipment,
    GiroEquipmentSnapshot,
    GiroImportStatus,
    GiroMonthlySale,
    GiroMonthlyTarget,
)
from app.models.pickup_catalog import (
    PickupCatalogClient,
    PickupCatalogInventoryItem,
    PickupCatalogUploadBatch,
)
from app.models.user import User
from app.schemas.giro import (
    GiroBreakdownOut,
    GiroImportResultOut,
    GiroImportStatusListOut,
    GiroImportStatusOut,
    GiroOverviewOut,
    GiroQuarterSummaryOut,
    GiroReportItemOut,
    GiroReportOut,
    GiroSummaryOut,
)
from app.services.giro_csv import (
    GiroCsvError,
    parse_sales_rows,
    parse_target_rows,
)
from app.services.giro_reference import EQUIPMENT_TYPE_BY_PRODUCT_CODE
from app.services.pickup_catalog_csv import canonical_code, parse_giro_date

router = APIRouter(prefix="/giro", tags=["Giro"])
get_giro_viewer = require_any_permission("giro.view", "giro.manage")
get_giro_manager = require_permission("giro.manage")

BRAZIL_TZ = ZoneInfo("America/Sao_Paulo")
INSTALLATION_CUTOFF = date(2023, 1, 1)
MONTHLY_TARGETS = {"visa": Decimal("1200.00"), "sopi": Decimal("2000.00")}
CENT = Decimal("0.01")
MAX_SALES_CSV_BYTES = 1024 * 1024 * 1024
MAX_MASTER_CSV_BYTES = 20 * 1024 * 1024
ALLOWED_UPLOAD_SUFFIXES = {".csv", ".txt"}
MESA_6_SECTORS = {"501", "502", "601", "602", "603", "605", "606"}
MESA_5_SECTORS = {"503", "504", "505", "506", "507", "604"}


def _month_start(value: date) -> date:
    return value.replace(day=1)


def _shift_month(value: date, offset: int) -> date:
    month_index = value.year * 12 + value.month - 1 + offset
    return date(month_index // 12, month_index % 12 + 1, 1)


def _current_month() -> str:
    return _month_start(datetime.now(BRAZIL_TZ).date()).strftime("%Y-%m")


def _report_months(selected_month: str | None = None) -> list[str]:
    current = date.fromisoformat(f"{selected_month or _current_month()}-01")
    return [_shift_month(current, offset).strftime("%Y-%m") for offset in (-3, -2, -1, 0)]


def mesa_for_sector(sector: str) -> str:
    normalized = str(sector or "").strip()
    if normalized in MESA_6_SECTORS:
        return "Mesa 6"
    if normalized in MESA_5_SECTORS:
        return "Mesa 5"
    return "Outros"


def _money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def _as_float(value: Decimal) -> float:
    return float(_money(value))


def _validate_upload_size(upload: UploadFile, maximum_bytes: int, label: str) -> int:
    filename = str(upload.filename or "").strip()
    if Path(filename).suffix.lower() not in ALLOWED_UPLOAD_SUFFIXES:
        raise HTTPException(status_code=400, detail=f"Envie o arquivo de {label} em formato CSV.")
    upload.file.seek(0, 2)
    size = upload.file.tell()
    upload.file.seek(0)
    if size > maximum_bytes:
        limit_gb = maximum_bytes / (1024 * 1024 * 1024)
        limit_mb = maximum_bytes / (1024 * 1024)
        limit_label = f"{limit_gb:g} GB" if limit_gb >= 1 else f"{limit_mb:g} MB"
        raise HTTPException(status_code=413, detail=f"O CSV de {label} excede o limite de {limit_label}.")
    if size == 0:
        raise HTTPException(status_code=400, detail=f"O CSV de {label} está vazio.")
    return size


async def _read_upload(upload: UploadFile, maximum_bytes: int, label: str) -> bytes:
    size = _validate_upload_size(upload, maximum_bytes, label)
    content = await upload.read(size + 1)
    return content


def _upsert_import_status(
    db: Session,
    dataset: str,
    filename: str,
    imported: int,
    ignored: int,
) -> GiroImportStatus:
    row = db.query(GiroImportStatus).filter(GiroImportStatus.dataset == dataset).first()
    if row is None:
        row = GiroImportStatus(dataset=dataset)
        db.add(row)
    row.file_name = Path(filename or "").name[:255]
    row.rows_imported = imported
    row.rows_ignored = ignored
    return row


def _commit_import(db: Session, *, dataset: str, filename: str, imported: int, ignored: int) -> GiroImportResultOut:
    record = _upsert_import_status(db, dataset, filename, imported, ignored)
    try:
        db.commit()
        db.refresh(record)
    except Exception:
        db.rollback()
        raise
    return GiroImportResultOut(
        message="Base importada com sucesso.",
        dataset=dataset,
        rows_imported=imported,
        rows_ignored=ignored,
        imported_at=record.imported_at,
    )


@router.post("/imports/sales", response_model=GiroImportResultOut)
async def import_sales(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_manager),
):
    del current_user
    _validate_upload_size(file, MAX_SALES_CSV_BYTES, "vendas")
    try:
        totals, source_months, imported, ignored = parse_sales_rows(file.file)
    except GiroCsvError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    months = sorted(source_months)
    try:
        db.query(GiroMonthlySale).filter(GiroMonthlySale.month.in_(months)).delete(synchronize_session=False)
        db.add_all(
            GiroMonthlySale(client_code=code, month=month, basket=basket, amount=_money(amount))
            for (code, month, basket), amount in totals.items()
        )
        return _commit_import(
            db,
            dataset="sales",
            filename=file.filename or "",
            imported=imported,
            ignored=ignored,
        )
    except Exception:
        db.rollback()
        raise


@router.post("/imports/targets", response_model=GiroImportResultOut)
async def import_targets(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_manager),
):
    del current_user
    raw = await _read_upload(file, MAX_MASTER_CSV_BYTES, "metas")
    try:
        targets = parse_target_rows(raw)
    except GiroCsvError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    months = sorted({target["month"] for target in targets})
    try:
        db.query(GiroMonthlyTarget).filter(GiroMonthlyTarget.month.in_(months)).delete(
            synchronize_session=False
        )
        db.add_all(GiroMonthlyTarget(**target) for target in targets)
        return _commit_import(
            db,
            dataset="targets",
            filename=file.filename or "",
            imported=len(targets),
            ignored=0,
        )
    except Exception:
        db.rollback()
        raise


def _import_statuses(db: Session) -> list[GiroImportStatus]:
    return db.query(GiroImportStatus).order_by(GiroImportStatus.dataset).all()


def _ensure_ready(db: Session, selected_month: str | None = None) -> None:
    completed = {row.dataset for row in _import_statuses(db)}
    missing = sorted({"sales", "targets"} - completed)
    if missing:
        raise HTTPException(
            status_code=409,
            detail="Importe a base anual de vendas (03.02.37) e as metas na área Atualizar base.",
        )
    if not db.query(PickupCatalogClient.id).first():
        raise HTTPException(status_code=409, detail="A base de clientes 01.20.11 não contém registros no banco de dados.")
    equipment_available = any(
        _equipment_for_month(db, equipment_type, selected_month or _current_month())[0]
        for equipment_type in MONTHLY_TARGETS
    )
    if not equipment_available:
        raise HTTPException(
            status_code=409,
            detail="A base de equipamentos 02.02.20 não contém comodatos VISA ou SOPI elegíveis no banco de dados.",
        )


def _selected_month(db: Session, requested_month: str | None) -> str:
    selected_month = requested_month or _current_month()
    if selected_month > _current_month():
        raise HTTPException(status_code=422, detail="A competência não pode ser futura.")
    return selected_month


def _equipment_by_client(db: Session, equipment_type: str) -> dict[str, Decimal]:
    result: dict[str, Decimal] = {}
    equipment_base = db.query(PickupCatalogInventoryItem)
    latest_equipment_batch = (
        db.query(PickupCatalogInventoryItem.batch_id)
        .filter(PickupCatalogInventoryItem.batch_id.isnot(None))
        .order_by(PickupCatalogInventoryItem.batch_id.desc())
        .first()
    )
    if latest_equipment_batch:
        equipment_base = equipment_base.filter(
            PickupCatalogInventoryItem.batch_id == latest_equipment_batch[0]
        )
    records = (
        equipment_base.join(
            PickupCatalogClient,
            PickupCatalogClient.id == PickupCatalogInventoryItem.client_id,
        )
        .with_entities(
            PickupCatalogClient.client_code,
            PickupCatalogInventoryItem.giro_equipment_type,
            PickupCatalogInventoryItem.product_code,
            PickupCatalogInventoryItem.giro_install_date,
            PickupCatalogInventoryItem.invoice_issue_date,
            PickupCatalogInventoryItem.giro_is_refrigerator,
            PickupCatalogInventoryItem.giro_balance,
            PickupCatalogInventoryItem.open_quantity,
            PickupCatalogInventoryItem.source_baixados,
        )
        .all()
    )
    for (
        client_code,
        stored_equipment_type,
        product_code,
        stored_install_date,
        issue_date,
        is_refrigerator,
        stored_balance,
        open_quantity,
        source_baixados,
    ) in records:
        mapped_type = EQUIPMENT_TYPE_BY_PRODUCT_CODE.get(canonical_code(product_code), "")
        row_equipment_type = str(stored_equipment_type or "").strip().lower()
        if row_equipment_type not in MONTHLY_TARGETS:
            row_equipment_type = mapped_type
        if row_equipment_type != equipment_type:
            continue

        install_date = stored_install_date or parse_giro_date(issue_date)
        if not install_date or install_date < INSTALLATION_CUTOFF:
            continue
        if not bool(is_refrigerator) and not mapped_type:
            continue

        balance = int(stored_balance or 0)
        if balance <= 0:
            balance = int(open_quantity or 0)
        if balance <= 0:
            balance = abs(int(source_baixados or 0))
        if balance <= 0:
            continue

        code = canonical_code(client_code)
        if code:
            result[code] = result.get(code, Decimal(0)) + Decimal(balance)
    return result


def _equipment_snapshot_by_client(
    db: Session,
    equipment_type: str,
    month: str,
) -> dict[str, Decimal]:
    result: dict[str, Decimal] = {}
    rows = (
        db.query(GiroEquipment.client_code, GiroEquipment.quantity)
        .filter(
            GiroEquipment.snapshot_month == month,
            GiroEquipment.equipment_type == equipment_type,
            GiroEquipment.install_date >= INSTALLATION_CUTOFF,
            GiroEquipment.is_refrigerator == 1,
            GiroEquipment.balance > 0,
        )
        .all()
    )
    for client_code, quantity in rows:
        code = canonical_code(client_code)
        if code:
            result[code] = result.get(code, Decimal(0)) + Decimal(str(quantity))
    return result


def _equipment_for_month(
    db: Session,
    equipment_type: str,
    selected_month: str,
) -> tuple[dict[str, Decimal], str]:
    current_month = _current_month()
    if selected_month == current_month:
        return _equipment_by_client(db, equipment_type), current_month

    snapshot_month = (
        db.query(GiroEquipmentSnapshot.month)
        .filter(GiroEquipmentSnapshot.month <= selected_month)
        .order_by(GiroEquipmentSnapshot.month.desc())
        .first()
    )
    if snapshot_month:
        return _equipment_snapshot_by_client(db, equipment_type, snapshot_month[0]), snapshot_month[0]

    return _equipment_by_client(db, equipment_type), current_month


def _report_rows(
    db: Session,
    equipment_type: str,
    months: list[str],
    *,
    equipment_month: str | None = None,
    sector: str = "",
    city: str = "",
    mesa: str = "",
    search: str = "",
) -> list[GiroReportItemOut]:
    equipment_counts = _equipment_for_month(
        db,
        equipment_type,
        equipment_month or _current_month(),
    )[0]
    codes = set(equipment_counts)
    if not codes:
        return []
    clients = {
        canonical_code(client.client_code): client
        for client in db.query(PickupCatalogClient)
        .filter(PickupCatalogClient.client_code.in_(codes))
        .all()
    }
    sales: dict[tuple[str, str], Decimal] = {}
    for code, month, amount in (
        db.query(GiroMonthlySale.client_code, GiroMonthlySale.month, GiroMonthlySale.amount)
        .filter(
            GiroMonthlySale.client_code.in_(codes),
            GiroMonthlySale.month.in_(months),
            GiroMonthlySale.basket == equipment_type,
        )
        .all()
    ):
        sales[(canonical_code(code), month)] = Decimal(amount)
    purchase_months: dict[str, str] = {}
    for client_code, month in (
        db.query(GiroMonthlySale.client_code, GiroMonthlySale.month)
        .filter(
            GiroMonthlySale.client_code.in_(codes),
            GiroMonthlySale.basket == equipment_type,
            GiroMonthlySale.amount > 0,
        )
        .order_by(GiroMonthlySale.month.desc())
        .all()
    ):
        purchase_months.setdefault(canonical_code(client_code), month)

    normalized_search = str(search or "").strip().lower()
    rows: list[GiroReportItemOut] = []
    target_per_equipment = MONTHLY_TARGETS[equipment_type]
    for code, count in equipment_counts.items():
        client = clients.get(code)
        client_sector = client.setor if client else ""
        client_city = client.cidade if client else ""
        mesa_name = mesa_for_sector(client_sector)
        if sector and client_sector != sector:
            continue
        if city and client_city != city:
            continue
        if mesa and mesa_name != mesa:
            continue
        fantasy = client.nome_fantasia if client else ""
        document = client.cnpj_cpf if client else ""
        status_text = client.status if client else ""
        frequency = client.frequency if client else ""
        searchable = " ".join((code, fantasy, document, client_sector, client_city)).lower()
        if normalized_search and normalized_search not in searchable:
            continue
        month_sales = {month: _as_float(sales.get((code, month), Decimal(0))) for month in months}
        current_sales = Decimal(str(month_sales[months[-1]]))
        monthly_target = _money(count * target_per_equipment)
        meets = current_sales >= monthly_target
        gap = max(Decimal(0), _money(monthly_target - current_sales))
        rows.append(
            GiroReportItemOut(
                client_code=code,
                fantasy_name=fantasy,
                document=document,
                client_status=status_text,
                frequency=frequency,
                sector=client_sector,
                city=client_city,
                mesa=mesa_name,
                equipment_count=float(count),
                month_sales=month_sales,
                last_purchase_month=purchase_months.get(code),
                monthly_target=_as_float(monthly_target),
                gap=_as_float(gap),
                giro_status="Atingindo" if meets else "Não atingiu",
            )
        )
    return sorted(rows, key=lambda item: (-item.gap, item.client_code))


def _summary(
    rows: list[GiroReportItemOut],
    current_month: str,
    target_percent: float | None = None,
) -> GiroSummaryOut:
    equipment_count = sum(Decimal(str(item.equipment_count)) for item in rows)
    ok_equipment = sum(
        (Decimal(str(item.equipment_count)) for item in rows if item.giro_status == "Atingindo"),
        Decimal(0),
    )
    monthly_target = sum((Decimal(str(item.monthly_target)) for item in rows), Decimal(0))
    current_sales = sum((Decimal(str(item.month_sales[current_month])) for item in rows), Decimal(0))
    gap = sum((Decimal(str(item.gap)) for item in rows), Decimal(0))
    percent = (ok_equipment / equipment_count * 100) if equipment_count else Decimal(0)
    return GiroSummaryOut(
        clients=len(rows),
        clients_not_meeting=sum(item.giro_status != "Atingindo" for item in rows),
        equipment_count=float(equipment_count),
        giro_ok_equipment=float(ok_equipment),
        giro_nok_equipment=float(equipment_count - ok_equipment),
        giro_ok_percent=round(float(percent), 2),
        target_percent=target_percent,
        monthly_target=_as_float(monthly_target),
        current_sales=_as_float(current_sales),
        gap=_as_float(gap),
    )


def _quarter_summary(
    db: Session,
    equipment_type: str,
    months: list[str],
    *,
    sector: str = "",
    city: str = "",
    mesa: str = "",
) -> GiroQuarterSummaryOut:
    current_month = months[-1]
    current_date = datetime.strptime(current_month, "%Y-%m").date()
    quarter_start = ((current_date.month - 1) // 3) * 3 + 1
    quarter_months = [
        f"{current_date.year:04d}-{month:02d}"
        for month in range(quarter_start, quarter_start + 3)
    ]
    elapsed_months = [month for month in quarter_months if month <= current_month]
    target_rows = (
        db.query(GiroMonthlyTarget)
        .filter(
            GiroMonthlyTarget.month.in_(quarter_months),
            GiroMonthlyTarget.equipment_type == equipment_type,
        )
        .all()
    )
    targets = {row.month: Decimal(row.target_percent) for row in target_rows}
    missing_targets = [month for month in quarter_months if month not in targets]
    target_percent = None
    if not missing_targets:
        target_percent = round(float(sum(targets.values(), Decimal(0)) / len(quarter_months)), 2)

    actual_current_month = _current_month()
    historical_months = [month for month in elapsed_months if month != actual_current_month]
    snapshot_months = {
        row.month
        for row in db.query(GiroEquipmentSnapshot.month)
        .filter(GiroEquipmentSnapshot.month.in_(historical_months))
        .all()
    } if historical_months else set()
    missing_snapshots = [month for month in historical_months if month not in snapshot_months]
    real_percent = None
    if not missing_snapshots:
        equipment_rows = (
            db.query(GiroEquipment)
            .filter(
                GiroEquipment.snapshot_month.in_(historical_months),
                GiroEquipment.equipment_type == equipment_type,
                GiroEquipment.install_date >= INSTALLATION_CUTOFF,
                GiroEquipment.is_refrigerator == 1,
                GiroEquipment.balance > 0,
            )
            .all()
        ) if historical_months else []
        current_equipment = (
            _equipment_by_client(db, equipment_type)
            if actual_current_month in elapsed_months else {}
        )
        codes = {row.client_code for row in equipment_rows} | set(current_equipment)
        clients = {
            canonical_code(client.client_code): client
            for client in db.query(PickupCatalogClient)
            .filter(PickupCatalogClient.client_code.in_(codes))
            .all()
        } if codes else {}
        equipment_by_month: dict[str, dict[str, Decimal]] = {}
        for equipment in equipment_rows:
            client = clients.get(canonical_code(equipment.client_code))
            client_sector = client.setor if client else ""
            client_city = client.cidade if client else ""
            if sector and client_sector != sector:
                continue
            if city and client_city != city:
                continue
            if mesa and mesa_for_sector(client_sector) != mesa:
                continue
            month_counts = equipment_by_month.setdefault(equipment.snapshot_month, {})
            code = canonical_code(equipment.client_code)
            month_counts[code] = month_counts.get(code, Decimal(0)) + Decimal(str(equipment.quantity))

        for code, count in current_equipment.items():
            client = clients.get(code)
            client_sector = client.setor if client else ""
            client_city = client.cidade if client else ""
            if sector and client_sector != sector:
                continue
            if city and client_city != city:
                continue
            if mesa and mesa_for_sector(client_sector) != mesa:
                continue
            equipment_by_month.setdefault(actual_current_month, {})[code] = count

        sales = {
            (code, month): Decimal(amount)
            for code, month, amount in (
                db.query(GiroMonthlySale.client_code, GiroMonthlySale.month, GiroMonthlySale.amount)
                .filter(
                    GiroMonthlySale.month.in_(elapsed_months),
                    GiroMonthlySale.basket == equipment_type,
                )
                .all()
            )
        }
        total_equipment = Decimal(0)
        giro_ok_equipment = Decimal(0)
        target_per_equipment = MONTHLY_TARGETS[equipment_type]
        for month in elapsed_months:
            for code, count in equipment_by_month.get(month, {}).items():
                total_equipment += count
                if sales.get((code, month), Decimal(0)) > count * target_per_equipment:
                    giro_ok_equipment += count
        if total_equipment:
            real_percent = round(float(giro_ok_equipment / total_equipment * 100), 2)

    return GiroQuarterSummaryOut(
        months=quarter_months,
        target_percent=target_percent,
        real_percent=real_percent,
        missing_target_months=missing_targets,
        missing_equipment_months=missing_snapshots,
    )


def _filters(db: Session) -> tuple[list[str], list[str]]:
    sectors = sorted({
        str(sector or "").strip()
        for (sector,) in db.query(PickupCatalogClient.setor).distinct().all()
        if str(sector or "").strip()
    })
    cities = sorted({
        str(city or "").strip()
        for (city,) in db.query(PickupCatalogClient.cidade).distinct().all()
        if str(city or "").strip()
    })
    return sectors, cities


def _target_percent(db: Session, equipment_type: str, month: str) -> float | None:
    target = (
        db.query(GiroMonthlyTarget.target_percent)
        .filter(
            GiroMonthlyTarget.month == month,
            GiroMonthlyTarget.equipment_type == equipment_type,
        )
        .first()
    )
    return round(float(target[0]), 2) if target else None


@router.get("/imports", response_model=GiroImportStatusListOut)
def get_import_status(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_viewer),
):
    del current_user
    sectors, cities = _filters(db)
    imports = [
        GiroImportStatusOut(
            dataset=item.dataset,
            file_name=item.file_name,
            rows_imported=item.rows_imported,
            rows_ignored=item.rows_ignored,
            imported_at=item.imported_at,
        )
        for item in _import_statuses(db)
    ]
    latest_base = db.query(PickupCatalogUploadBatch).order_by(PickupCatalogUploadBatch.id.desc()).first()
    if latest_base is not None:
        snapshot_status = (
            db.query(GiroEquipmentSnapshot)
            .order_by(GiroEquipmentSnapshot.month.desc())
            .first()
        )
        imports.extend([
            GiroImportStatusOut(
                dataset="clients",
                file_name=latest_base.clients_file_name or "",
                rows_imported=int(latest_base.clients_count or 0),
                rows_ignored=0,
                imported_at=latest_base.uploaded_at,
            ),
            GiroImportStatusOut(
                dataset="equipment",
                file_name=(
                    snapshot_status.file_name
                    if snapshot_status else latest_base.inventory_file_name or ""
                ),
                rows_imported=(
                    snapshot_status.rows_imported
                    if snapshot_status else int(latest_base.open_items or 0)
                ),
                rows_ignored=snapshot_status.rows_ignored if snapshot_status else 0,
                imported_at=snapshot_status.imported_at if snapshot_status else latest_base.uploaded_at,
            ),
        ])
    snapshot_months = [
        row.month
        for row in db.query(GiroEquipmentSnapshot.month).order_by(GiroEquipmentSnapshot.month.desc()).all()
    ]
    sales_by_equipment: dict[str, set[str]] = {"visa": set(), "sopi": set()}
    for month, equipment_type in (
        db.query(GiroMonthlySale.month, GiroMonthlySale.basket).distinct().all()
    ):
        if equipment_type in sales_by_equipment:
            sales_by_equipment[equipment_type].add(month)
    sales_months = set.union(*sales_by_equipment.values())
    current_month = _current_month()
    available_months = sorted(
        (set(snapshot_months) | sales_months | {current_month}) - {
            month
            for month in set(snapshot_months) | sales_months
            if month > current_month
        },
        reverse=True,
    )
    months_without_snapshot = sorted(
        {month for month in sales_months if month < current_month} - set(snapshot_months),
        reverse=True,
    )
    target_years = sorted({
        int(row.month[:4])
        for row in db.query(GiroMonthlyTarget.month).distinct().all()
    })
    return GiroImportStatusListOut(
        imports=imports,
        sectors=sectors,
        cities=cities,
        equipment_snapshot_months=snapshot_months,
        available_months=available_months,
        months_without_equipment_snapshot=months_without_snapshot,
        sales_months_by_equipment={
            equipment_type: sorted(months, reverse=True)
            for equipment_type, months in sales_by_equipment.items()
        },
        target_years=target_years,
    )


@router.get("/reports/{equipment_type}", response_model=GiroReportOut)
def get_report(
    equipment_type: str,
    month: str | None = Query(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
    sector: str = Query(default="", max_length=80),
    city: str = Query(default="", max_length=120),
    mesa: str = Query(default="", pattern="^(|Mesa 5|Mesa 6|Outros)$"),
    search: str = Query(default="", max_length=120),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_viewer),
):
    del current_user
    if equipment_type not in MONTHLY_TARGETS:
        raise HTTPException(status_code=404, detail="Tipo de equipamento inválido.")
    selected_month = _selected_month(db, month)
    _ensure_ready(db, selected_month)
    months = _report_months(selected_month)
    equipment_reference_month = _equipment_for_month(db, equipment_type, selected_month)[1]
    rows = _report_rows(
        db,
        equipment_type,
        months,
        equipment_month=selected_month,
        sector=sector,
        city=city,
        mesa=mesa,
        search=search,
    )
    sectors, cities = _filters(db)
    start = (page - 1) * page_size
    return GiroReportOut(
        equipment_type=equipment_type,
        months=months,
        equipment_reference_month=equipment_reference_month,
        equipment_is_estimated=equipment_reference_month != selected_month,
        summary=_summary(rows, selected_month, _target_percent(db, equipment_type, selected_month)),
        tri=_quarter_summary(
            db,
            equipment_type,
            months,
            sector=sector,
            city=city,
            mesa=mesa,
        ),
        items=rows[start:start + page_size],
        total_items=len(rows),
        page=page,
        page_size=page_size,
        sectors=sectors,
        cities=cities,
    )


def _breakdown(
    visa_rows: list[GiroReportItemOut],
    sopi_rows: list[GiroReportItemOut],
    current_month: str,
    key: str,
    target_per_type: dict[str, float | None],
) -> list[GiroBreakdownOut]:
    rows = visa_rows + sopi_rows
    names = sorted({getattr(item, key) or "Sem informação" for item in rows})
    result: list[GiroBreakdownOut] = []
    for name in names:
        visa_group = [item for item in visa_rows if (getattr(item, key) or "Sem informação") == name]
        sopi_group = [item for item in sopi_rows if (getattr(item, key) or "Sem informação") == name]
        result.append(
            GiroBreakdownOut(
                name=name,
                mesa=mesa_for_sector(name) if key == "sector" else None,
                visa=_summary(visa_group, current_month, target_per_type["visa"]),
                sopi=_summary(sopi_group, current_month, target_per_type["sopi"]),
            )
        )
    return result


@router.get("/overview", response_model=GiroOverviewOut)
def get_overview(
    month: str | None = Query(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
    city: str = Query(default="", max_length=120),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_viewer),
):
    del current_user
    selected_month = _selected_month(db, month)
    _ensure_ready(db, selected_month)
    months = _report_months(selected_month)
    equipment_references = {
        equipment_type: _equipment_for_month(db, equipment_type, selected_month)[1]
        for equipment_type in MONTHLY_TARGETS
    }
    equipment_reference_month = min(equipment_references.values(), default=selected_month)
    visa = _report_rows(db, "visa", months, equipment_month=selected_month, city=city)
    sopi = _report_rows(db, "sopi", months, equipment_month=selected_month, city=city)
    target_per_type = {
        equipment_type: _target_percent(db, equipment_type, selected_month)
        for equipment_type in MONTHLY_TARGETS
    }
    return GiroOverviewOut(
        month=selected_month,
        equipment_reference_month=equipment_reference_month,
        equipment_is_estimated=equipment_reference_month != selected_month,
        visa_equipment_reference_month=equipment_references["visa"],
        visa_equipment_is_estimated=equipment_references["visa"] != selected_month,
        sopi_equipment_reference_month=equipment_references["sopi"],
        sopi_equipment_is_estimated=equipment_references["sopi"] != selected_month,
        visa=_summary(visa, selected_month, target_per_type["visa"]),
        sopi=_summary(sopi, selected_month, target_per_type["sopi"]),
        visa_tri=_quarter_summary(db, "visa", months, city=city),
        sopi_tri=_quarter_summary(db, "sopi", months, city=city),
        by_sector=_breakdown(visa, sopi, selected_month, "sector", target_per_type),
        by_mesa=_breakdown(visa, sopi, selected_month, "mesa", target_per_type),
        by_city=_breakdown(visa, sopi, selected_month, "city", target_per_type),
    )


def _make_workbook(equipment_type: str, months: list[str], items: list[GiroReportItemOut]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"GIRO {equipment_type.upper()}"
    headers = [
        "Código do PDV", "Fantasia", "Documento", "Status", "Frequência", "Quantidade de equipamento",
        *[f"Faturamento {month}" for month in months],
        "Último mês com compra", "Meta do PDV", "GAP", "Status do Giro", "Setor", "Cidade", "Mesa",
    ]
    sheet.append(headers)
    for item in items:
        sheet.append([
            item.client_code,
            item.fantasy_name,
            item.document,
            item.client_status,
            item.frequency,
            item.equipment_count,
            *[item.month_sales[month] for month in months],
            item.last_purchase_month or "",
            item.monthly_target,
            item.gap,
            item.giro_status,
            item.sector,
            item.city,
            item.mesa,
        ])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2, min_col=7, max_col=6 + len(months)):
        for cell in row:
            cell.number_format = '"R$" #,##0.00'
    for row in sheet.iter_rows(min_row=2, min_col=8 + len(months), max_col=9 + len(months)):
        for cell in row:
            cell.number_format = '"R$" #,##0.00'
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(32, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/reports/{equipment_type}/export")
def export_report(
    equipment_type: str,
    month: str | None = Query(default=None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
    sector: str = Query(default="", max_length=80),
    city: str = Query(default="", max_length=120),
    mesa: str = Query(default="", pattern="^(|Mesa 5|Mesa 6|Outros)$"),
    search: str = Query(default="", max_length=120),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_giro_viewer),
):
    del current_user
    if equipment_type not in MONTHLY_TARGETS:
        raise HTTPException(status_code=404, detail="Tipo de equipamento inválido.")
    selected_month = _selected_month(db, month)
    _ensure_ready(db, selected_month)
    months = _report_months(selected_month)
    rows = _report_rows(
        db,
        equipment_type,
        months,
        equipment_month=selected_month,
        sector=sector,
        city=city,
        mesa=mesa,
        search=search,
    )
    workbook_file = _make_workbook(equipment_type, months, rows)
    filename = f"giro-{equipment_type}-{selected_month}.xlsx"
    return StreamingResponse(
        workbook_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
