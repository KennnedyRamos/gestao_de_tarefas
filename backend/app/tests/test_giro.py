import asyncio
from io import BytesIO
from datetime import date
from decimal import Decimal

from fastapi import HTTPException, UploadFile
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook

from app.database.base import Base
from app.models.giro import (
    GiroEquipment,
    GiroEquipmentSnapshot,
    GiroImportStatus,
    GiroMonthlySale,
    GiroMonthlyTarget,
)
from app.models.pickup_catalog import PickupCatalogClient, PickupCatalogInventoryItem
from app.routes.giro import (
    _ensure_ready,
    _equipment_by_client,
    _make_workbook,
    _quarter_summary,
    _report_rows,
    _selected_month,
    _summary,
    _validate_upload_size,
    import_sales,
    MAX_SALES_CSV_BYTES,
    mesa_for_sector,
)
from app.services.giro_csv import (
    parse_sales_rows,
    parse_target_rows,
)
from app.services.giro_reference import BASKET_BY_PRODUCT_CODE, EQUIPMENT_TYPE_BY_PRODUCT_CODE
from app.services.pickup_catalog_csv import load_clients_csv, load_inventory_csv


def test_existing_client_csv_import_includes_status_and_frequency_for_giro():
    csv_data = (
        "Codigo Cliente;Nome Fantasia;CNPJ;Status do Cliente;Frequencia Visita;Cod. Setor;Cidade\n"
        "100;Bar Central;12345678000199;A;Semanal;501;Registro\n"
    ).encode("utf-8")

    result = load_clients_csv(csv_data)

    client = result["100"]
    assert client["nome_fantasia"] == "Bar Central"
    assert client["cnpj_cpf"] == "12345678000199"
    assert client["status"] == "A"
    assert client["frequency"] == "Semanal"
    assert client["setor"] == "501"
    assert client["cidade"] == "Registro"


def test_existing_inventory_csv_import_captures_giro_equipment_fields():
    csv_data = (
        "Categoria;Codigo Cliente;Descricao;Saldo;Equipamento;Data Operacao\n"
        "SOPI;100;SOPI VISA COOLER;-2;Sim;02/10/2026\n"
        "GFA 1L;100;GFA 1L;-3;Não;02/10/2026\n"
    ).encode("utf-8")

    result = load_inventory_csv(csv_data)

    assert len(result["100"]) == 2
    refrigerator = result["100"][0]
    assert refrigerator["giro_equipment_type"] == "sopi"
    assert refrigerator["giro_is_refrigerator"] is True
    assert refrigerator["giro_install_date"] == date(2026, 10, 2)
    assert refrigerator["giro_balance"] == 2


def test_inventory_csv_uses_issue_date_when_operation_date_is_invalid():
    csv_data = (
        "Codigo Cliente;Descricao;Saldo;Codigo Produto;Data Operacao;Data Emissao\n"
        "100;REFRIGERADOR SKOL;-1;0118780;00/00/0000;10/09/2026\n"
    ).encode("utf-8")

    item = load_inventory_csv(csv_data)["100"][0]

    assert item["giro_equipment_type"] == "sopi"
    assert item["giro_is_refrigerator"] is True
    assert item["giro_install_date"] == date(2026, 9, 10)


def test_reference_files_classify_sales_and_equipment_products():
    assert len(BASKET_BY_PRODUCT_CODE) == 2167
    assert BASKET_BY_PRODUCT_CODE["132"] == "sopi"
    assert BASKET_BY_PRODUCT_CODE["80"] == "visa"
    assert BASKET_BY_PRODUCT_CODE["13859"] == "sopi"
    assert EQUIPMENT_TYPE_BY_PRODUCT_CODE["118780"] == "sopi"
    assert EQUIPMENT_TYPE_BY_PRODUCT_CODE["862277"] == "visa"

    inventory_csv = (
        "Codigo Cliente;Descricao;Saldo;Codigo Produto;Data Operacao\n"
        "100;Expositor vertical; -2;118780;02/10/2026\n"
    ).encode("utf-8")
    item = load_inventory_csv(inventory_csv)["100"][0]
    assert item["giro_equipment_type"] == "sopi"
    assert item["giro_is_refrigerator"] is True


def test_giro_uses_equipment_mapping_for_legacy_inventory_rows():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    client = PickupCatalogClient(client_code="00100", nome_fantasia="Bar Central")
    session.add(client)
    session.add_all([
        GiroImportStatus(dataset="sales", file_name="sales.csv"),
        GiroImportStatus(dataset="targets", file_name="targets.csv"),
    ])
    session.flush()

    def add_inventory(
        product_code: str,
        *,
        issue_date: str = "02/10/2026",
        balance: int = -2,
    ) -> None:
        quantity = abs(balance)
        session.add(PickupCatalogInventoryItem(
            client_id=client.id,
            description="Comodato",
            item_type="outro",
            product_code=product_code,
            invoice_issue_date=issue_date,
            source_baixados=balance,
            open_quantity=quantity,
            giro_equipment_type="",
            giro_install_date=None,
            giro_is_refrigerator=False,
            giro_balance=0,
        ))

    add_inventory("118724")
    add_inventory("118780")
    add_inventory("188005")
    add_inventory("118724", issue_date="31/12/2022")
    add_inventory("118780", balance=0)
    session.commit()

    _ensure_ready(session)
    assert _equipment_by_client(session, "visa") == {"100": Decimal(2)}
    assert _equipment_by_client(session, "sopi") == {"100": Decimal(2)}
    session.close()
    engine.dispose()


def test_sales_csv_excludes_marketplace_and_chopp_and_aggregates_baskets():
    csv_data = (
        "NAB;CERVEJA;PDV;Emissao;Total;Origem do Pedido;Descrição;Status;Indicador Cancelamento\n"
        "NAB;-;100;02/10/2026;1.200,00;B2BG;GUARANA LATA;A;N\n"
        "-;CTT;100;02/10/2026;2.000,00;B2BG;SKOL LATA;A;N\n"
        "NAB;-;100;02/10/2026;9.000,00;Marketplace;GUARANA LATA;A;N\n"
        "-;CTT;100;02/10/2026;9.000,00;B2BG;CHOPP BIB;A;N\n"
    ).encode("utf-8")

    totals, source_months, imported, ignored = parse_sales_rows(csv_data)

    assert totals[("100", "2026-10", "visa")] == Decimal("1200.00")
    assert totals[("100", "2026-10", "sopi")] == Decimal("2000.00")
    assert source_months == {"2026-10"}
    assert imported == 2
    assert ignored == 2


def test_sales_csv_uses_product_reference_when_baskets_are_not_in_sales_file():
    csv_data = (
        "PDV;Emissao;Valor Total;Origem do Pedido;Cod Produto;Descricao\n"
        "100;02/10/2026;1.500,00;B2BG;80;REFRIGERANTE LATA\n"
        "100;02/10/2026;2.500,00;B2BG;13859;CERVEJA LATA\n"
    ).encode("utf-8")

    totals, source_months, imported, ignored = parse_sales_rows(csv_data)

    assert totals[("100", "2026-10", "visa")] == Decimal("1500.00")
    assert totals[("100", "2026-10", "sopi")] == Decimal("2500.00")
    assert source_months == {"2026-10"}
    assert imported == 2
    assert ignored == 0


def test_sales_csv_file_stream_is_parsed_without_loading_into_bytes():
    csv_data = (
        "PDV;Emissao;Total;Origem;Codigo Produto\n"
        "100;02/10/2026;1.500,00;B2BG;80\n"
    ).encode("utf-8")

    totals, _, imported, ignored = parse_sales_rows(BytesIO(csv_data))

    assert totals[("100", "2026-10", "visa")] == Decimal("1500.00")
    assert imported == 1
    assert ignored == 0


def test_sales_csv_accepts_real_product_header_and_mixed_ansi_bytes():
    csv_data = (
        "Cliente;Emissao;Produto;Total;Origem do Pedido;Descri\x90o;Status\n"
        "100;02/10/2026;80;1.200,00;B2BG;REFRIGERANTE LATA;A\n"
        "100;02/10/2026;132;2.500,00;B2BG;CERVEJA LATA;A\n"
        "100;02/10/2026;80;9.000,00;B2BG;CHOPP BIB;A\n"
    ).encode("latin-1")

    totals, source_months, imported, ignored = parse_sales_rows(BytesIO(csv_data))

    assert totals[("100", "2026-10", "visa")] == Decimal("1200.00")
    assert totals[("100", "2026-10", "sopi")] == Decimal("2500.00")
    assert source_months == {"2026-10"}
    assert imported == 2
    assert ignored == 1


def test_sales_csv_uses_client_column_c_and_operation_date_column_f():
    headers = [f"Extra {column}" for column in range(1, 65)]
    headers[2] = "Cliente .."
    headers[5] = "Dt. Operacao"
    headers[6] = "Emissao"
    headers[9] = "Status"
    headers[12] = "Cliente"
    headers[15] = "Produto"
    headers[17] = "Descricao"
    headers[25] = "Desconto"
    headers[26] = "Total"
    headers[63] = "Origem do Pedido"
    row = [""] * len(headers)
    row[2] = "00100"
    row[5] = "30/09/2026"
    row[6] = "01/10/2026"
    row[9] = "A"
    row[12] = "99999"
    row[15] = "80"
    row[17] = "REFRIGERANTE LATA"
    row[25] = "500,00"
    row[26] = "1.200,00"
    row[63] = "B2BG"
    csv_data = (
        ";".join(headers) + "\n" + ";".join(row) + "\n"
    ).encode("utf-8")

    totals, source_months, imported, ignored = parse_sales_rows(csv_data)

    assert totals[("100", "2026-09", "visa")] == Decimal("1200.00")
    assert source_months == {"2026-09"}
    assert imported == 1
    assert ignored == 0


def test_sales_upload_limit_accepts_files_over_200_mb_and_rejects_above_500_mb():
    class SizedFile:
        def __init__(self, size):
            self.size = size
            self.position = 0

        def seek(self, offset, whence=0):
            self.position = self.size if whence == 2 else offset

        def tell(self):
            return self.position

    over_200_mb = UploadFile(
        filename="sales.csv",
        file=SizedFile(201 * 1024 * 1024),
    )
    assert _validate_upload_size(over_200_mb, MAX_SALES_CSV_BYTES, "vendas") == 201 * 1024 * 1024

    over_limit = UploadFile(
        filename="sales.csv",
        file=SizedFile(MAX_SALES_CSV_BYTES + 1),
    )
    try:
        _validate_upload_size(over_limit, MAX_SALES_CSV_BYTES, "vendas")
    except HTTPException as exc:
        assert exc.status_code == 413
    else:
        raise AssertionError("Expected oversized sales upload to be rejected.")


def test_sales_csv_normalizes_client_codes_and_keeps_months_with_no_eligible_sales():
    csv_data = (
        "NAB;CERVEJA;MATCH;PDV;Emissao;Total;Origem do Pedido;Descrição;Status\n"
        "NAB;-;-;00100;01/09/2026;1.200,00;B2BG;GUARANA LATA;A\n"
        "-;-;MATCH;00100;02/09/2026;2.000,00;B2BG;MATCH LATA;A\n"
        "NAB;-;-;100;02/10/2026;9.000,00;Marketplace;GUARANA LATA;A\n"
    ).encode("utf-8")

    totals, source_months, imported, ignored = parse_sales_rows(csv_data)

    assert totals[("100", "2026-09", "visa")] == Decimal("1200.00")
    assert totals[("100", "2026-09", "sopi")] == Decimal("2000.00")
    assert source_months == {"2026-09", "2026-10"}
    assert imported == 2
    assert ignored == 1


def test_sales_csv_allows_a_month_with_only_excluded_sales():
    csv_data = (
        "NAB;CERVEJA;PDV;Emissao;Total;Origem do Pedido;Descrição;Status\n"
        "NAB;-;100;01/09/2026;1.200,00;Marketplace;GUARANA LATA;A\n"
    ).encode("utf-8")

    totals, source_months, imported, ignored = parse_sales_rows(csv_data)

    assert totals == {}
    assert source_months == {"2026-09"}
    assert imported == 0
    assert ignored == 1


def test_sales_import_removes_previous_month_values_when_new_file_has_only_excluded_sales():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add(GiroMonthlySale(
        client_code="100",
        month="2026-09",
        basket="visa",
        amount=Decimal("1200.00"),
    ))
    session.commit()
    csv_data = (
        "NAB;CERVEJA;PDV;Emissao;Total;Origem do Pedido;Descrição;Status\n"
        "NAB;-;100;01/09/2026;1.200,00;Marketplace;GUARANA LATA;A\n"
    ).encode("utf-8")

    result = asyncio.run(import_sales(
        file=UploadFile(filename="vendas.csv", file=BytesIO(csv_data)),
        db=session,
        current_user=None,
    ))

    assert result.rows_imported == 0
    assert session.query(GiroMonthlySale).count() == 0
    status = session.query(GiroImportStatus).one()
    assert status.dataset == "sales"
    assert status.rows_imported == 0
    session.close()
    engine.dispose()


def test_targets_csv_reads_percentages_for_visa_and_sopi():
    csv_data = (
        "Indicador;Ano;Jan;Fev;Mar\n"
        "GIRO VISA;2026;0.09;0.085;0.075\n"
        "GIRO SOPI;2026;0.73;0.628;0.621\n"
    ).encode("utf-8")

    targets = parse_target_rows(csv_data)

    values = {(target["month"], target["equipment_type"]): target["target_percent"] for target in targets}
    assert values[("2026-01", "visa")] == Decimal("9.00")
    assert values[("2026-02", "visa")] == Decimal("8.500")
    assert values[("2026-01", "sopi")] == Decimal("73.00")


def test_giro_report_counts_exact_target_as_meeting_and_counts_all_client_equipment():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    client = PickupCatalogClient(
        client_code="100",
        nome_fantasia="Bar Central",
        cnpj_cpf="123",
        status="A",
        frequency="Semanal",
        setor="501",
        cidade="Registro",
    )
    session.add(client)
    session.flush()
    session.add(PickupCatalogInventoryItem(
        client_id=client.id,
        description="VISA Cooler",
        item_type="refrigerador",
        giro_equipment_type="visa",
        giro_install_date=date(2023, 1, 1),
        giro_is_refrigerator=True,
        giro_balance=2,
    ))
    session.add_all([
        GiroImportStatus(dataset="sales", file_name="sales.csv"),
        GiroImportStatus(dataset="targets", file_name="targets.csv"),
    ])
    session.add(GiroMonthlySale(
        client_code="100",
        month="2026-10",
        basket="visa",
        amount=Decimal("2400.00"),
    ))
    session.commit()

    _ensure_ready(session)
    months = ["2026-07", "2026-08", "2026-09", "2026-10"]
    rows = _report_rows(session, "visa", months)
    summary = _summary(rows, months[-1], target_percent=8.5)

    assert rows[0].equipment_count == 2
    assert rows[0].monthly_target == 2400
    assert rows[0].giro_status == "Atingindo"
    assert rows[0].gap == 0
    assert summary.clients_not_meeting == 0
    assert summary.giro_ok_percent == 100
    assert summary.giro_nok_equipment == 0
    assert summary.target_percent == 8.5

    sale = session.query(GiroMonthlySale).one()
    sale.amount = Decimal("2399.99")
    session.commit()
    below_target_row = _report_rows(session, "visa", months)[0]
    assert below_target_row.giro_status == "Não atingiu"
    assert below_target_row.gap == 0.01
    assert mesa_for_sector("601") == "Mesa 6"
    assert mesa_for_sector("604") == "Mesa 5"
    assert mesa_for_sector("11") == "Outros"

    workbook = load_workbook(_make_workbook("visa", months, rows), read_only=True)
    sheet = workbook.active
    assert sheet.cell(1, 1).value == "Código do PDV"
    assert sheet.cell(1, 12).value == "Meta do PDV"
    assert sheet.cell(2, 1).value == "100"
    assert sheet.cell(2, 13).value == 0

    workbook.close()
    session.close()
    engine.dispose()


def test_historical_report_uses_snapshot_and_includes_last_purchase_month():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    client = PickupCatalogClient(client_code="100", nome_fantasia="Bar Central", setor="501")
    session.add(client)
    session.flush()
    session.add(PickupCatalogInventoryItem(
        client_id=client.id,
        description="VISA Cooler atual",
        giro_equipment_type="visa",
        giro_install_date=date(2023, 1, 1),
        giro_is_refrigerator=True,
        giro_balance=7,
    ))
    session.add(GiroEquipmentSnapshot(month="2024-01"))
    session.add(GiroEquipment(
        client_code="100",
        equipment_type="visa",
        snapshot_month="2024-01",
        install_date=date(2023, 1, 1),
        quantity=2,
        balance=2,
        is_refrigerator=True,
    ))
    session.add_all([
        GiroMonthlySale(client_code="100", month="2023-12", basket="visa", amount=Decimal("50.00")),
        GiroMonthlySale(client_code="100", month="2024-01", basket="visa", amount=Decimal("100.00")),
    ])
    session.commit()

    assert _selected_month(session, "2024-01") == "2024-01"
    rows = _report_rows(
        session,
        "visa",
        ["2023-10", "2023-11", "2023-12", "2024-01"],
        equipment_month="2024-01",
    )

    assert rows[0].equipment_count == 2
    assert rows[0].month_sales["2024-01"] == 100
    assert rows[0].last_purchase_month == "2024-01"
    session.close()
    engine.dispose()


def test_quarter_meta_is_simple_mean_and_real_is_weighted_by_equipment():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    client = PickupCatalogClient(
        client_code="100",
        nome_fantasia="Bar Central",
        cnpj_cpf="123",
        status="A",
        frequency="Semanal",
        setor="501",
        cidade="Registro",
    )
    session.add(client)
    session.flush()
    months = ["2026-07", "2026-08", "2026-09"]
    counts = [Decimal("1"), Decimal("2"), Decimal("1")]
    sales = [Decimal("1200.01"), Decimal("2400.00"), Decimal("1200.01")]
    metas = [Decimal("9"), Decimal("8.5"), Decimal("7.5")]
    for month, count, amount, target in zip(months, counts, sales, metas):
        session.add(GiroEquipmentSnapshot(month=month, file_name=f"{month}.csv"))
        session.add(GiroEquipment(
            client_code="100",
            equipment_type="visa",
            snapshot_month=month,
            install_date=date(2023, 1, 1),
            quantity=count,
            balance=count,
            is_refrigerator=True,
        ))
        session.add(GiroMonthlySale(
            client_code="100",
            month=month,
            basket="visa",
            amount=amount,
        ))
        session.add(GiroMonthlyTarget(
            month=month,
            equipment_type="visa",
            target_percent=target,
        ))
    session.add(PickupCatalogInventoryItem(
        client_id=client.id,
        description="VISA Cooler",
        item_type="refrigerador",
        giro_equipment_type="visa",
        giro_install_date=date(2023, 1, 1),
        giro_is_refrigerator=True,
        giro_balance=1,
    ))
    session.commit()

    quarter = _quarter_summary(session, "visa", ["2026-06", *months])

    assert quarter.months == months
    assert quarter.target_percent == 8.33
    assert quarter.real_percent == 50
    assert quarter.missing_equipment_months == []
    assert quarter.missing_target_months == []
    session.close()
    engine.dispose()
